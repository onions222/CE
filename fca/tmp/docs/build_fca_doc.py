from pathlib import Path
import numpy as np
from PIL import Image
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/onion/Desktop/code/CE/fca')
TMP = ROOT / 'tmp/docs'
OUT = ROOT / 'output/doc'
TMP.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

orig_path = ROOT / 'soap-bubbles-nature.jpg'
out_path = ROOT / 'out_hue_strict_vs.png'
flowchart_path = TMP / 'fca_flowchart.png'
compare_path = TMP / 'fca_compare.png'
docx_path = OUT / 'FCA_算法说明文档.docx'


def set_east_asia_font(style, latin='Times New Roman', east_asia='Microsoft YaHei', size=11):
    style.font.name = latin
    style.font.size = Pt(size)
    rpr = style.element.rPr
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), east_asia)


def add_bold_par(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    return p


# ---------- 1) simplified flowchart ----------
fig, ax = plt.subplots(figsize=(10, 11))
ax.set_xlim(0, 10)
ax.set_ylim(0, 14)
ax.axis('off')

boxes = [
    (1.2, 12.2, 7.6, 1.1, 'Start: read RGB image and parameters\n(deltaH, hue range, taper_width)'),
    (1.2, 10.4, 7.6, 1.1, 'Compute per-pixel V, Delta, S and original Hue\nby 6-sector piecewise mapping'),
    (1.2, 8.6, 7.6, 1.1, 'Build active mask in hue range\n(optional taper near range boundaries)'),
    (1.2, 6.8, 7.6, 1.1, 'Apply hue shift: new_hue = (Hue + eff_dH) mod 360\nand map to new_sector + new_t'),
    (1.2, 5.0, 7.6, 1.1, 'Reconstruct RGB with fixed V and Delta\n=> preserve S = Delta / V'),
    (1.2, 3.2, 7.6, 1.1, 'YCoCg round-trip interface + clip [0,1]\nOutput image and error metrics'),
]
for x, y, w, h, text in boxes:
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle='round,pad=0.02,rounding_size=0.08',
        linewidth=1.5, edgecolor='#1f2937', facecolor='#eef6ff'
    )
    ax.add_patch(patch)
    ax.text(x + w/2, y + h/2, text, ha='center', va='center', fontsize=9.5)

for i in range(len(boxes)-1):
    x1, y1, w1, h1, _ = boxes[i]
    x2, y2, w2, h2, _ = boxes[i+1]
    ax.annotate('', xy=(x2+w2/2, y2+h2), xytext=(x1+w1/2, y1), arrowprops=dict(arrowstyle='->', lw=1.8, color='#2563eb'))

ax.set_title('FCA.py Flowchart (Simplified)', fontsize=15, pad=14)
fig.tight_layout()
fig.savefig(flowchart_path, dpi=220)
plt.close(fig)

# ---------- 2) redraw comparison image with exactly two requested files ----------
orig = np.asarray(Image.open(orig_path).convert('RGB'))
out = np.asarray(Image.open(out_path).convert('RGB'))

fig2, axes = plt.subplots(1, 2, figsize=(15, 6.5))
axes[0].imshow(orig)
axes[0].set_title('soap-bubbles-nature.jpg (Original)', fontsize=12)
axes[0].axis('off')
axes[1].imshow(out)
axes[1].set_title('out_hue_strict_vs.png (Processed)', fontsize=12)
axes[1].axis('off')
fig2.suptitle('Before vs After (Exact Source Files)', fontsize=15)
fig2.tight_layout()
fig2.savefig(compare_path, dpi=220)
plt.close(fig2)

# ---------- 3) build docx ----------
doc = Document()

# style tuning
normal = doc.styles['Normal']
set_east_asia_font(normal, latin='Times New Roman', east_asia='Microsoft YaHei', size=11)

# heading styles
for lv in ['Heading 1', 'Heading 2']:
    st = doc.styles[lv]
    set_east_asia_font(st, latin='Calibri', east_asia='Microsoft YaHei', size=14 if lv == 'Heading 1' else 12)

h = doc.add_heading('FCA.py 算法说明文档', level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('版本: 基于当前 FCA.py 实现（cross-sector hue shift + hue range taper）')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1 background
doc.add_heading('1. 算法开发背景', level=1)
doc.add_paragraph(
    '本算法用于图像色相编辑场景。核心需求是：在调整色相(Hue)时，尽量保持亮度(Value)和饱和度(Saturation)稳定，'
    '并支持只在指定色相带内进行局部编辑。'
)
doc.add_paragraph(
    '相比“直接在 RGB 空间旋转色度”方法，本实现采用分段 HSV 语义重建，目的是让颜色变化更可控、可解释。'
)

# 2 objective
doc.add_heading('2. 算法目标', level=1)
for t in [
    '支持跨扇区色相位移: H -> H + deltaH。',
    '在重建中保持 V 与 Delta 不变，从而保持 S = Delta / V。',
    '支持色相范围掩码与边缘平滑(taper)，减轻范围边界突变。',
]:
    doc.add_paragraph(t, style='List Bullet')

# 3 formulas (academic style)
doc.add_heading('3. 详细实现与数学公式（学术化表述）', level=1)
doc.add_paragraph('本章采用与 FCA_fixedpoint.py 对齐的变量体系，并给出编号公式、结论与误差项。')

add_bold_par(doc, '3.1 预备定义与有效域')
doc.add_paragraph('设输入像素为 x=(R,G,B), R,G,B∈[0,1]。定义如下:')
doc.add_paragraph('V(x)=max(R,G,B),   m(x)=min(R,G,B),   Delta(x)=V(x)-m(x)    (3-1)')
doc.add_paragraph('S(x)=0 (V=0), 否则 S(x)=Delta(x)/V(x)                           (3-2)')
doc.add_paragraph('当 Delta=0 时 Hue 无定义，此类像素记为非活动像素并保持原值。')

add_bold_par(doc, '3.2 扇区化 Hue 表示')
doc.add_paragraph('活动像素按 (max通道, min通道) 划分为 6 个扇区 s∈{0,1,2,3,4,5}。')
doc.add_paragraph('在每个扇区定义线性参数 t∈[0,1]，例如:')
doc.add_paragraph('t=(G-B)/Delta (s=0),  t=(G-R)/Delta (s=1),  t=(B-R)/Delta (s=2)   (3-3)')
doc.add_paragraph('t=(B-G)/Delta (s=3),  t=(R-G)/Delta (s=4),  t=(R-B)/Delta (s=5)   (3-4)')
doc.add_paragraph('Hue 可写为:')
doc.add_paragraph('H = 60*(s+t),   H∈[0,360)                                           (3-5)')
doc.add_paragraph('对应 fixed-point 形式为 hue_hu = s*ONE + t_fp，与 FCA_fixedpoint.py 一致。')

add_bold_par(doc, '3.3 色相范围与回绕判定')
doc.add_paragraph('给定区间 [hmin,hmax]，活动掩码 A 定义为:')
doc.add_paragraph('A={H|hmin<=H<=hmax}, 若 hmin<=hmax                                  (3-6)')
doc.add_paragraph('A={H|H>=hmin 或 H<=hmax}, 若 hmin>hmax                              (3-7)')
doc.add_paragraph('式(3-7)处理跨 0° 区间（如 [300°,30°]）。')

add_bold_par(doc, '3.4 taper 平滑与有效位移')
doc.add_paragraph('记 d(H) 为 H 到区间边界的最小角距离，定义:')
doc.add_paragraph('scale(H)=clip(d(H)/taper_width, 0, 1)                               (3-8)')
doc.add_paragraph('eff_dH(H)=deltaH*scale(H)                                           (3-9)')
doc.add_paragraph('当 taper_width=0 时，退化为 eff_dH=deltaH。')

add_bold_par(doc, '3.5 Hue 更新与新扇区定位')
doc.add_paragraph('更新后的色相为:')
doc.add_paragraph('H_new = (H + eff_dH) mod 360                                        (3-10)')
doc.add_paragraph('s_new = floor(H_new/60),   t_new = H_new/60 - s_new                (3-11)')
doc.add_paragraph('fixed-point 对应关系: s_new = hue_new_hu >> frac_bits, t_new = hue_new_hu & (ONE-1)。')

add_bold_par(doc, '3.6 通道重建')
doc.add_paragraph('在 s_new 扇区中按“max/min固定 + 中间通道插值”重建:')
doc.add_paragraph('c_max=V,   c_min=m=V-Delta                                           (3-12)')
doc.add_paragraph('c_mid=c_min+Delta*t_new 或 c_mid=c_max-Delta*t_new                  (3-13)')
doc.add_paragraph('具体采用哪一式由扇区定义决定，且与 FCA_fixedpoint.py 的 ns0..ns5 规则一一对应。')

add_bold_par(doc, '3.7 定理与结论')
doc.add_paragraph('定理 1（亮度保持）: 对所有活动像素，重建后 V_out = V_in。')
doc.add_paragraph('证明要点: 式(3-12)直接将最大通道赋值为 V，且其它通道不超过该值。')
doc.add_paragraph('定理 2（饱和度保持）: 对所有活动像素，S_out = S_in。')
doc.add_paragraph('证明要点: 由式(3-12)得 Delta_out = c_max-c_min = V-(V-Delta)=Delta；再代入式(3-2)即得。')
doc.add_paragraph('推论: 本算法的主要变化仅体现在 Hue 维度，V 与 S 在数学上保持不变。')

add_bold_par(doc, '3.8 误差项说明')
doc.add_paragraph('理论上式(3-1)至式(3-13)在实数域可精确满足约束；工程实现误差来自以下三类:')
doc.add_paragraph('1) 浮点舍入误差: 由除法和乘法产生，量级通常在 1e-15 至 1e-12。')
doc.add_paragraph('2) 量化误差: 保存到 uint8 时引入 ±0.5 LSB 误差。')
doc.add_paragraph('3) 定点近似误差: fixed-point 中由 frac_bits 限定，单次乘法截断误差上界约为 2^(-frac_bits)。')
doc.add_paragraph('本文档采用的验证指标为:')
doc.add_paragraph('maxErrV = max|V_out - V_in|,   maxErrS = max|S_out - S_in|         (3-14)')
doc.add_paragraph('并可附加 MAE/MaxE(R,G,B) 作为与参考浮点实现的一致性指标。')

# 4 flowchart
doc.add_heading('4. 标准流程图（简化版）', level=1)
doc.add_paragraph('流程图如下，合并了细枝末节，保留核心处理主线。')
doc.add_picture(str(flowchart_path), width=Inches(6.3))

# 5 comparison
doc.add_heading('5. 算法作用后对比图', level=1)
doc.add_paragraph('以下对比图严格使用两个指定文件: soap-bubbles-nature.jpg 与 out_hue_strict_vs.png。')
doc.add_picture(str(compare_path), width=Inches(6.6))

# 6 notes
doc.add_heading('6. 工程建议', level=1)
for t in [
    '建议设置 |deltaH| 的上限阈值，避免过大位移带来非期望颜色迁移。',
    '建议保持 taper_width >= deltaH_max，以进一步抑制范围边界不连续。',
    '低饱和度区域可增加 Delta 最小阈值，减少噪声触发。',
]:
    doc.add_paragraph(t, style='List Bullet')

# 7 fixed-point hardware resource model
doc.add_heading('7. 定点实现所需硬件资源（参考 FCA_fixedpoint.py）', level=1)
doc.add_paragraph('本节给出“按算法结构推导”的资源模型，用于硬件规划。以下为计算单元需求，不包含片上总线与外设资源。')

res_table = doc.add_table(rows=1, cols=4)
res_table.style = 'Table Grid'
h0 = res_table.rows[0].cells
h0[0].text = '资源类型'
h0[1].text = '主要用途'
h0[2].text = '典型位宽'
h0[3].text = '并行/复用需求'

rows = [
    ('比较器', '计算 V/m 与扇区判定 (R>=G 等)', '8-10 bit', '每像素需多路比较; 流水可复用'),
    ('加法器/减法器', 'Delta, 通道差分, 重建加减', '9-18 bit', '主路径必需, 数量与并行像素数线性相关'),
    ('乘法器', 'Delta*t_new 与 taper 缩放', '约 (8+frac_bits) bit', '可 1-2 个乘法器时分复用'),
    ('除法器', 't=diff/Delta 与 scale=d/taper', '被除数约 (8+frac_bits) bit', '可用迭代除法器复用降低面积'),
    ('移位器', '>> frac_bits 与 << frac_bits', '与数据通路同宽', '固定移位, 成本较低'),
    ('逻辑单元', '掩码组合与扇区多路选择', '1 bit 掩码网络', '与像素并行度线性增长'),
    ('寄存器', '流水级缓存与中间量保持', '若干 8-18 bit 寄存器组', '建议按级插入以满足时序'),
]
for r in rows:
    c = res_table.add_row().cells
    c[0].text, c[1].text, c[2].text, c[3].text = r

doc.add_paragraph('存储资源建议:')
doc.add_paragraph('1) 行缓存/帧缓存: 若逐像素流式处理，可不需要全帧缓存，仅需输入输出行缓冲。', style='List Number')
doc.add_paragraph('2) 中间变量寄存: 需保留 V, m, Delta, hue_hu/new_hue, new_sec, new_t 等。', style='List Number')
doc.add_paragraph('3) 参数寄存: hmin/hmax, deltaH_hu, taper_hu, frac_bits, delta_eps。', style='List Number')

doc.add_paragraph('面积-性能估算公式（以每时钟处理 P 个像素计）:')
doc.add_paragraph('比较器数量 ~ O(P), 加减器数量 ~ O(P), 乘法器与除法器数量取决于是否时分复用。')
doc.add_paragraph('若采用单通道复用除法器，吞吐降低但面积显著下降；若全并行，则吞吐最大但 DSP/逻辑占用更高。')

doc.save(docx_path)
print(docx_path)
print(flowchart_path)
print(compare_path)
